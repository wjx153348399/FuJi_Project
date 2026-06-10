import csv
import datetime
import hashlib
import json
import os
import re
import subprocess
import time
import zipfile
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from xml.etree import ElementTree

import requests
from requests import exceptions as requests_exceptions
from zoneinfo import ZoneInfo

try:
    import openpyxl
except ImportError:  # pragma: no cover - optional dependency
    openpyxl = None


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(BASE_DIR, "config.json")

with open(CONFIG_FILE, "r", encoding="utf-8") as f:
    config = json.load(f)

API_LOGIN_URL = config["api"]["login_url"]
API_UPLOAD_URL = config["api"]["upload_url"]
API_USERNAME = config["api"]["username"]
API_PASSWORD = config["api"]["password"]

SCAN_ROOT = config["paths"]["scan_root"]
LOG_DIR = config["paths"]["log_root"]

SCAN_RECURSIVE = config["scan"].get("recursive", True)
ALLOWED_EXTENSIONS = tuple(ext.lower() for ext in config["scan"].get("extensions", [".xls", ".xlsx"]))
EXCLUDE_PREFIXES = tuple(config["scan"].get("exclude_prefixes", ["~$"]))
EXCLUDE_DIRS = set(config["scan"].get("exclude_dirs", []))

TIMEZONE = ZoneInfo(config["schedule"].get("timezone", "Asia/Shanghai"))
DAY_OFFSET = int(config["schedule"].get("day_offset", 1))
EXECUTE_TIME = config["schedule"].get("execute_time", "00:10")

DEDUPE_MODE = config["dedupe"].get("mode", "business_key_first")
KEEP_LATEST_IN_BATCH = config["dedupe"].get("keep_latest_in_batch", True)

MAX_RETRIES = int(config["upload"].get("max_retries", 1))
RETRY_DELAY = int(config["upload"].get("retry_delay_seconds", 3))
UPLOAD_TIMEOUT = int(config["upload"].get("timeout_seconds", 60))
RETRY_ON_TIMEOUT = bool(config["upload"].get("retry_on_timeout", False))
SERIAL_NO_MODE = config["upload"].get("serial_no_mode", "empty")
SERIAL_NO_PREFIX = config["upload"].get("serial_no_prefix", "ZK")
MAX_WORKERS = int(config["upload"].get("max_workers", 5))

ENABLE_PENDING_CSV = bool(config["pending"].get("enable_csv", True))

BUSINESS_RECORDS_FILE = os.path.join(LOG_DIR, "uploaded_business_keys.json")
FINGERPRINT_RECORDS_FILE = os.path.join(LOG_DIR, "uploaded_file_fingerprints.json")
PENDING_TIMEOUT_FILE = os.path.join(LOG_DIR, "pending_timeout.json")

F26_PATTERN = re.compile(r"(F\d{2}-\d{3}-\d{4})", re.IGNORECASE)

login_lock = threading.Lock()


def ensure_log_dir():
    os.makedirs(LOG_DIR, exist_ok=True)


def now_local():
    return datetime.datetime.now(TIMEZONE)


def format_dt(dt_obj):
    return dt_obj.astimezone(TIMEZONE).strftime("%Y-%m-%d %H:%M:%S")


def load_json_dict(file_path):
    ensure_log_dir()
    if not os.path.exists(file_path):
        return {}
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, OSError):
        return {}


def load_json_list(file_path):
    ensure_log_dir()
    if not os.path.exists(file_path):
        return []
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
            return data if isinstance(data, list) else []
    except (json.JSONDecodeError, OSError):
        return []


def save_json(file_path, data):
    ensure_log_dir()
    temp_file = f"{file_path}.tmp"
    with open(temp_file, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(temp_file, file_path)


def append_jsonl(file_path, entry):
    ensure_log_dir()
    with open(file_path, "a", encoding="utf-8") as f:
        f.write(json.dumps(entry, ensure_ascii=False) + "\n")


def generate_run_id():
    return f"{SERIAL_NO_PREFIX}-{now_local().strftime('%Y%m%d-%H%M%S')}"


def get_upload_serial_no():
    if SERIAL_NO_MODE == "empty":
        return ""
    if SERIAL_NO_MODE == "run_batch":
        return generate_run_id()
    return ""


def calculate_file_hash(file_path):
    sha256_hash = hashlib.sha256()
    with open(file_path, "rb") as f:
        for byte_block in iter(lambda: f.read(1024 * 1024), b""):
            sha256_hash.update(byte_block)
    return sha256_hash.hexdigest()


def normalize_filename_stem(file_name):
    stem = os.path.splitext(file_name)[0].upper()
    stem = stem.replace("（", "(").replace("）", ")")
    stem = re.sub(r"\s+", "-", stem)
    stem = stem.replace("_", "-")
    stem = re.sub(r"-{2,}", "-", stem)
    return stem.strip("-")


def extract_f26_code(file_name):
    match = F26_PATTERN.search(file_name)
    return match.group(1).upper() if match else None


def extract_business_key(file_name):
    normalized_stem = normalize_filename_stem(file_name)
    f26_code = extract_f26_code(normalized_stem)
    if not f26_code:
        return None

    body = normalized_stem.split(f26_code, 1)[0].strip("-_")
    body = re.sub(r"-{2,}", "-", body).strip("-")
    if not body:
        return None
    return f"{body}|{f26_code}"


def get_target_day_window():
    current = now_local()
    target_day = (current - datetime.timedelta(days=DAY_OFFSET)).date()
    start_dt = datetime.datetime.combine(target_day, datetime.time.min, tzinfo=TIMEZONE)
    end_dt = start_dt + datetime.timedelta(days=1)
    return target_day, start_dt, end_dt


def mount_share_drive():
    host = config["share_drive"]["host"]
    user = config["share_drive"]["username"]
    pwd = config["share_drive"]["password"]

    if os.path.exists(host):
        return

    print(f"尝试挂载共享盘: {host}")
    cmd = f'net use "{host}" "{pwd}" /user:"{user}"'
    try:
        subprocess.run(cmd, shell=True, check=True, capture_output=True)
        print("共享盘挂载成功")
    except subprocess.CalledProcessError as e:
        raise RuntimeError(f"共享盘挂载失败: {e.stderr.decode(errors='ignore') or e}") from e


def login(session):
    payload = {
        "username": API_USERNAME,
        "password": API_PASSWORD
    }
    response = session.post(API_LOGIN_URL, json=payload, timeout=15)
    response.raise_for_status()
    result = response.json()
    if str(result.get("code")) != "20220":
        raise RuntimeError(f"登录失败: {result.get('message', '未知错误')}")

    token = result.get("token")
    if not token:
        raise RuntimeError("登录失败: 未获取到 token")

    session.headers.update({
        "Authorization": f"Bearer {token}"
    })
    return token


def scan_candidate_files(start_dt, end_dt):
    candidates = []
    if not os.path.exists(SCAN_ROOT):
        print(f"警告: 扫描根目录不存在 - {SCAN_ROOT}")
        return candidates

    for root, dirs, files in os.walk(SCAN_ROOT):
        dirs[:] = [dir_name for dir_name in dirs if dir_name not in EXCLUDE_DIRS]
        if not SCAN_RECURSIVE:
            dirs[:] = []

        for file_name in files:
            file_name_lower = file_name.lower()
            if file_name.startswith(EXCLUDE_PREFIXES):
                continue
            if not file_name_lower.endswith(ALLOWED_EXTENSIONS):
                continue

            file_path = os.path.join(root, file_name)
            if not os.path.isfile(file_path):
                continue

            mtime_ts = os.path.getmtime(file_path)
            mtime_dt = datetime.datetime.fromtimestamp(mtime_ts, TIMEZONE)
            if start_dt <= mtime_dt < end_dt:
                candidates.append(file_path)
    return candidates


def build_file_metadata(file_path):
    filename = os.path.basename(file_path)
    file_hash = calculate_file_hash(file_path)
    business_key = extract_business_key(filename) if DEDUPE_MODE == "business_key_first" else None
    normalized_stem = normalize_filename_stem(filename)
    fallback_key = f"FALLBACK|{normalized_stem}|{file_hash}"
    dedupe_key = f"BUSINESS|{business_key}" if business_key else fallback_key
    stat = os.stat(file_path)
    return {
        "filename": filename,
        "full_path": file_path,
        "directory": os.path.dirname(file_path),
        "file_hash": file_hash,
        "file_size": stat.st_size,
        "mtime_ts": stat.st_mtime,
        "file_mtime": format_dt(datetime.datetime.fromtimestamp(stat.st_mtime, TIMEZONE)),
        "business_key": business_key,
        "fallback_key": fallback_key,
        "dedupe_key": dedupe_key,
        "query_hint": build_query_hint(business_key)
    }


def build_query_hint(business_key):
    if not business_key:
        return "无法自动提取业务键，请按文件名和目录人工核对。"
    model, batch = business_key.split("|", 1)
    return f"建议在QMS按 料号={model}、批次号={batch} 查询。"


def choose_preferred_file(left_info, right_info):
    if not KEEP_LATEST_IN_BATCH:
        return left_info, right_info
    if right_info["mtime_ts"] > left_info["mtime_ts"]:
        return right_info, left_info
    return left_info, right_info


def dedupe_candidates(candidate_files):
    selected = {}
    batch_duplicates = []

    for file_path in candidate_files:
        info = build_file_metadata(file_path)
        existing = selected.get(info["dedupe_key"])
        if existing is None:
            selected[info["dedupe_key"]] = info
            continue

        keep_info, skip_info = choose_preferred_file(existing, info)
        selected[info["dedupe_key"]] = keep_info
        skip_info["skip_reason"] = "batch_duplicate_skipped"
        skip_info["message"] = (
            f"批内重复，仅保留较新文件。保留文件: {keep_info['full_path']}"
        )
        batch_duplicates.append(skip_info)

    return list(selected.values()), batch_duplicates


def load_success_records():
    return load_json_dict(BUSINESS_RECORDS_FILE), load_json_dict(FINGERPRINT_RECORDS_FILE)


def was_uploaded(info, business_records, fingerprint_records):
    if info["business_key"] and info["business_key"] in business_records:
        return True, "history_duplicate_skipped", "命中历史业务键去重记录"
    if info["fallback_key"] in fingerprint_records:
        return True, "history_duplicate_skipped", "命中历史文件指纹去重记录"
    return False, "", ""


def save_success_record(info, run_id, upload_result, business_records, fingerprint_records):
    record = {
        "filename": info["filename"],
        "full_path": info["full_path"],
        "directory": info["directory"],
        "file_hash": info["file_hash"],
        "file_size": info["file_size"],
        "file_mtime": info["file_mtime"],
        "run_id": run_id,
        "uploaded_at": format_dt(now_local()),
        "response": upload_result
    }
    fingerprint_records[info["fallback_key"]] = record
    if info["business_key"]:
        business_records[info["business_key"]] = record
    save_json(BUSINESS_RECORDS_FILE, business_records)
    save_json(FINGERPRINT_RECORDS_FILE, fingerprint_records)


def extract_excel_summary(file_path):
    summary = {
        "sheet_names": [],
        "row_count": None,
        "first_row_preview": None,
        "last_row_preview": None,
        "summary_note": ""
    }

    ext = os.path.splitext(file_path)[1].lower()
    if ext != ".xlsx":
        summary["summary_note"] = "仅对 .xlsx 做详细摘要，.xls 仅记录基础信息。"
        return summary

    if openpyxl is not None:
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            summary["sheet_names"] = workbook.sheetnames
            first_sheet = workbook[workbook.sheetnames[0]]
            row_count = 0
            first_row = None
            last_row = None
            for row in first_sheet.iter_rows(values_only=True):
                row_values = [value for value in row if value not in (None, "")]
                if not row_values:
                    continue
                row_count += 1
                row_preview = [str(value) for value in row[:8]]
                if first_row is None:
                    first_row = row_preview
                last_row = row_preview
            summary["row_count"] = row_count
            summary["first_row_preview"] = first_row
            summary["last_row_preview"] = last_row
            return summary
        except Exception as exc:  # pragma: no cover - best effort
            summary["summary_note"] = f"openpyxl 摘要失败: {exc}"

    try:
        with zipfile.ZipFile(file_path) as zf:
            workbook_xml = ElementTree.fromstring(zf.read("xl/workbook.xml"))
            ns = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            summary["sheet_names"] = [
                node.attrib.get("name", "")
                for node in workbook_xml.findall(".//a:sheets/a:sheet", ns)
            ]
            summary["summary_note"] = "未安装 openpyxl，仅提取到工作表名称。"
    except Exception as exc:  # pragma: no cover - best effort
        summary["summary_note"] = f"zip/xml 摘要失败: {exc}"
    return summary


def write_log(info, status, message, run_id, retry_count=0, response=None):
    log_file = os.path.join(LOG_DIR, f"upload_log_{now_local().date()}.jsonl")
    entry = {
        "filename": info["filename"],
        "full_path": info["full_path"],
        "directory": info["directory"],
        "file_hash": info["file_hash"],
        "file_size": info["file_size"],
        "file_mtime": info["file_mtime"],
        "business_key": info["business_key"],
        "fallback_key": info["fallback_key"],
        "dedupe_key": info["dedupe_key"],
        "run_id": run_id,
        "status": status,
        "message": message,
        "retry_count": retry_count,
        "log_time": format_dt(now_local()),
        "response": response
    }
    append_jsonl(log_file, entry)


def append_pending_confirmation(info, run_id, error_message):
    pending_records = load_json_list(PENDING_TIMEOUT_FILE)
    summary = extract_excel_summary(info["full_path"])
    record = {
        "status": "pending_confirm",
        "run_id": run_id,
        "timeout_at": format_dt(now_local()),
        "filename": info["filename"],
        "full_path": info["full_path"],
        "directory": info["directory"],
        "file_hash": info["file_hash"],
        "file_size": info["file_size"],
        "file_mtime": info["file_mtime"],
        "business_key": info["business_key"],
        "fallback_key": info["fallback_key"],
        "dedupe_key": info["dedupe_key"],
        "query_hint": info["query_hint"],
        "error_message": error_message,
        "excel_summary": summary
    }
    pending_records.append(record)
    save_json(PENDING_TIMEOUT_FILE, pending_records)
    if ENABLE_PENDING_CSV:
        write_pending_csv(pending_records)
    return record


def write_pending_csv(records):
    csv_file = os.path.join(LOG_DIR, f"pending_timeout_{now_local().date()}.csv")
    fieldnames = [
        "status", "run_id", "timeout_at", "filename", "full_path", "directory",
        "file_mtime", "file_size", "file_hash", "business_key", "dedupe_key",
        "query_hint", "error_message", "row_count", "sheet_names"
    ]
    with open(csv_file, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for record in records:
            excel_summary = record.get("excel_summary", {})
            writer.writerow({
                "status": record.get("status"),
                "run_id": record.get("run_id"),
                "timeout_at": record.get("timeout_at"),
                "filename": record.get("filename"),
                "full_path": record.get("full_path"),
                "directory": record.get("directory"),
                "file_mtime": record.get("file_mtime"),
                "file_size": record.get("file_size"),
                "file_hash": record.get("file_hash"),
                "business_key": record.get("business_key"),
                "dedupe_key": record.get("dedupe_key"),
                "query_hint": record.get("query_hint"),
                "error_message": record.get("error_message"),
                "row_count": excel_summary.get("row_count"),
                "sheet_names": ",".join(excel_summary.get("sheet_names", []))
            })


def write_summary(run_id, target_day, stats, failed_files, pending_files):
    summary_file = os.path.join(
        LOG_DIR,
        f"summary_{now_local().strftime('%Y-%m-%d_%H-%M-%S')}.json"
    )
    summary_data = {
        "run_id": run_id,
        "run_time": format_dt(now_local()),
        "schedule_execute_time": EXECUTE_TIME,
        "target_day": str(target_day),
        "stats": stats,
        "failed_files": failed_files,
        "pending_files": pending_files
    }
    save_json(summary_file, summary_data)


def upload_file(session, file_path, serial_no):
    filename = os.path.basename(file_path)
    attempts_allowed = MAX_RETRIES + 1

    for attempt in range(attempts_allowed):
        try:
            with open(file_path, "rb") as f:
                content_type = "application/vnd.ms-excel"
                if filename.lower().endswith(".xlsx"):
                    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                files = {"file": (filename, f, content_type)}
                data = {"serialNo": serial_no}
                response = session.post(API_UPLOAD_URL, files=files, data=data, timeout=UPLOAD_TIMEOUT)
                
                if response.status_code == 401:
                    print(f"[{filename}] 认证失效 (401)，正在重新登录... ({attempt + 1}/{MAX_RETRIES})")
                    try:
                        with login_lock:
                            login(session)
                    except Exception as login_exc:
                        print(f"[{filename}] 重新登录失败: {login_exc}")
                    if attempt < MAX_RETRIES:
                        time.sleep(RETRY_DELAY)
                        continue
                    else:
                        return "failed", {"message": f"401 Unauthorized: {response.text}"}, attempt

                response.raise_for_status()
                result = response.json()

                if (
                    result.get("success") is True
                    or result.get("code") in (200, "200", 20000, "20000")
                ):
                    return "success", result, attempt

                msg = result.get("message", "未知业务错误")
                if attempt < MAX_RETRIES:
                    print(f"[{filename}] 业务失败: {msg}，准备重试 ({attempt + 1}/{MAX_RETRIES})")
                    time.sleep(RETRY_DELAY)
                    continue
                return "failed", result, attempt

        except requests_exceptions.ReadTimeout as exc:
            if RETRY_ON_TIMEOUT and attempt < MAX_RETRIES:
                print(f"[{filename}] 读取超时，准备重试 ({attempt + 1}/{MAX_RETRIES})")
                time.sleep(RETRY_DELAY)
                continue
            return "pending_confirm", {"message": str(exc)}, attempt
        except requests_exceptions.Timeout as exc:
            if RETRY_ON_TIMEOUT and attempt < MAX_RETRIES:
                print(f"[{filename}] 请求超时，准备重试 ({attempt + 1}/{MAX_RETRIES})")
                time.sleep(RETRY_DELAY)
                continue
            return "pending_confirm", {"message": str(exc)}, attempt
        except Exception as exc:
            if attempt < MAX_RETRIES:
                print(f"[{filename}] 请求异常: {exc}，准备重试 ({attempt + 1}/{MAX_RETRIES})")
                time.sleep(RETRY_DELAY)
                continue
            return "failed", {"message": str(exc)}, attempt

    return "failed", {"message": "重试次数耗尽"}, MAX_RETRIES


def main():
    current_run_time = now_local()
    print(f"开始执行阻抗文件批量上传任务: {format_dt(current_run_time)}")
    mount_share_drive()

    run_id = generate_run_id()
    upload_serial_no = get_upload_serial_no()
    target_day, start_dt, end_dt = get_target_day_window()
    print(f"本次执行标识(run_id): {run_id}")
    print(f"计划执行时间: {EXECUTE_TIME}，本次处理日期: {target_day}")

    session = requests.Session()
    adapter = requests.adapters.HTTPAdapter(pool_connections=MAX_WORKERS, pool_maxsize=MAX_WORKERS)
    session.mount('http://', adapter)
    session.mount('https://', adapter)
    
    try:
        login(session)
        print("登录成功，开始准备上传。")
    except Exception as exc:
        print(f"登录失败，任务终止: {exc}")
        return

    candidate_files = scan_candidate_files(start_dt, end_dt)
    print(f"找到 {len(candidate_files)} 个前一天修改的 Excel 文件")
    if not candidate_files:
        return

    selected_files, batch_duplicates = dedupe_candidates(candidate_files)
    for duplicate in batch_duplicates:
        print(f"批内重复跳过: {duplicate['filename']} -> {duplicate['message']}")
        write_log(duplicate, duplicate["skip_reason"], duplicate["message"], run_id)

    business_records, fingerprint_records = load_success_records()

    success_count = 0
    fail_count = 0
    skip_count = len(batch_duplicates)
    pending_count = 0
    failed_files = []
    pending_files = []

    futures = {}
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        for info in selected_files:
            uploaded, skip_status, skip_message = was_uploaded(info, business_records, fingerprint_records)
            if uploaded:
                print(f"历史重复跳过: {info['filename']} ({info['query_hint']})")
                write_log(info, skip_status, skip_message, run_id)
                skip_count += 1
                continue

            print(f"加入上传队列: {info['filename']} (路径: {info['full_path']}) ...")
            future = executor.submit(upload_file, session, info["full_path"], upload_serial_no)
            futures[future] = info

        for future in as_completed(futures):
            info = futures[future]
            try:
                status, response, retries = future.result()
            except Exception as exc:
                status, response, retries = "failed", {"message": str(exc)}, 0

            message = response.get("message", "")

            if status == "success":
                print(f"上传成功: {info['filename']}")
                write_log(info, "success", message, run_id, retries, response)
                save_success_record(info, run_id, response, business_records, fingerprint_records)
                success_count += 1
            elif status == "pending_confirm":
                print(f"上传待确认: {info['filename']} - {message}")
                pending_record = append_pending_confirmation(info, run_id, message)
                write_log(info, "pending_confirm", message, run_id, retries, response)
                pending_files.append(pending_record)
                pending_count += 1
            else:
                print(f"上传失败: {info['filename']} - {message}")
                write_log(info, "failed", message, run_id, retries, response)
                failed_info = {
                    "filename": info["filename"],
                    "full_path": info["full_path"],
                    "business_key": info["business_key"],
                    "message": message
                }
                failed_files.append(failed_info)
                fail_count += 1

    stats = {
        "candidate_count": len(candidate_files),
        "selected_count": len(selected_files),
        "success_count": success_count,
        "fail_count": fail_count,
        "skip_count": skip_count,
        "pending_count": pending_count
    }
    write_summary(run_id, target_day, stats, failed_files, pending_files)
    print("-" * 30)
    print(
        f"任务结束。候选: {len(candidate_files)}, 入选: {len(selected_files)}, "
        f"成功: {success_count}, 失败: {fail_count}, 跳过: {skip_count}, 待确认: {pending_count}"
    )
    print("-" * 30)


if __name__ == "__main__":
    main()
